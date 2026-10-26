import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from datetime import datetime
from bson import ObjectId
from src.client import api, Cursor


def test_dashboard_grouping_office_selection_and_clear(qtbot, monkeypatch):
    from src.ui.dashboard import DashboardWidget, local_time, csv_safe
    kano, abuja = ObjectId(), ObjectId()
    def call(path, body=None):
        assert path == '/read'
        if body['collection'] == 'employees':
            return [{'office_id': 'engineering-kano', '_id': kano, 'employee_no': 'K1', 'full_name': 'Kano Staff', 'status': 'active'},
                    {'office_id': 'logistics-abuja', '_id': abuja, 'employee_no': 'A1', 'full_name': 'Abuja Staff', 'status': 'active'}]
        return [{'office_id': 'engineering-kano', 'employee_id': kano, 'work_date': '2026-10-19', 'clock_in_at': datetime(2026, 10, 19, 8)},
                {'office_id': 'logistics-abuja', 'employee_id': abuja, 'work_date': '2026-10-19', 'clock_in_at': datetime(2026, 10, 19, 7)}]
    monkeypatch.setattr(api, 'call', call)
    monkeypatch.setattr(api, 'token', 'test-token')
    dashboard = DashboardWidget()
    qtbot.addWidget(dashboard)
    dashboard.set_user({'role': 'main_admin'})
    assert dashboard.attendance_table.rowCount() == 3
    assert 'Date:' in dashboard.attendance_table.item(0, 0).text()
    assert dashboard.attendance_table.item(1, 0).text() == 'Abuja'
    assert dashboard.attendance_table.item(2, 0).text() == 'Kano'
    assert local_time(datetime(2026, 10, 19, 8)) == '09:00:00'
    assert csv_safe('=BAD()') == "'=BAD()"
    dashboard.clear_data()
    assert dashboard.attendance_table.rowCount() == dashboard.emp_table.rowCount() == 0
    dashboard.set_user({'role': 'office_admin', 'office_id': 'engineering-kano'})
    assert not dashboard.office_filter.isEnabled()
    assert dashboard.office_filter.currentData() == 'engineering-kano'


def test_main_window_starts_with_login_without_database(qtbot, monkeypatch):
    from src.scanner.provider import KojackVendorProvider
    monkeypatch.setattr(KojackVendorProvider, 'initialize', lambda self: False)
    monkeypatch.setattr(api, 'token', None)
    from main import MainWindow
    window = MainWindow()
    qtbot.addWidget(window)
    assert window.stacked_widget.currentWidget() is window.login_view
    assert window.dashboard_view.emp_table.rowCount() == 0


def test_legacy_conversion_preserves_templates_and_rejects_code():
    import io
    import pickle
    import zlib
    import numpy as np
    import mongomock
    import pytest
    from src.server.legacy import convert_templates, TemplateReader
    db = mongomock.MongoClient().test
    sample = np.ones((100, 32), dtype=np.uint8)
    emp = db.employees.insert_one({'fingerprint_descriptors': zlib.compress(pickle.dumps([sample]))}).inserted_id
    assert convert_templates(db) == (1, 0)
    assert db.employees.find_one({'_id': emp})['fingerprint_samples'] == [sample.tolist()]
    assert convert_templates(db) == (0, 0)
    with pytest.raises(ValueError):
        TemplateReader(io.BytesIO(b'cos\nsystem\n.')).load()


def test_dashboard_session_timeout_and_deletion_controls(qtbot, monkeypatch):
    from src.ui.dashboard import DashboardWidget, DeleteHistoryDialog
    kano = ObjectId()
    def call(path, body=None):
        if body['collection'] == 'employees':
            return [{'office_id': 'engineering-kano', '_id': kano, 'employee_no': 'K1', 'full_name': 'Kano Staff', 'status': 'active'}]
        return [{'office_id': 'engineering-kano', 'employee_id': kano, 'work_date': '2026-10-19', 'clock_in_at': datetime(2026, 10, 19, 8)}]
    monkeypatch.setattr(api, 'call', call)
    monkeypatch.setattr(api, 'token', 'test-token')

    dashboard = DashboardWidget()
    qtbot.addWidget(dashboard)

    # 1. Verify 10-minute session idle timeout (600,000 ms)
    assert dashboard.inactivity_timer.interval() == 600_000

    # 2. Master Admin controls
    dashboard.set_user({'role': 'main_admin'})
    dashboard.show()
    assert dashboard.inactivity_timer.isActive()
    assert dashboard.delete_history_btn.isVisible()
    assert dashboard.delete_selected_btn.isVisible()
    assert not dashboard.delete_selected_btn.isEnabled()

    # Selecting date header row (row 0) does not enable delete
    dashboard.attendance_table.selectRow(0)
    assert not dashboard.delete_selected_btn.isEnabled()

    # Selecting actual attendance record (row 1) enables the sleek Delete Selected Log button
    dashboard.attendance_table.selectRow(1)
    assert dashboard.delete_selected_btn.isEnabled()

    # Deselecting disables it
    dashboard.attendance_table.clearSelection()
    assert not dashboard.delete_selected_btn.isEnabled()

    # 3. Office Admin cannot see deletion controls
    dashboard.set_user({'role': 'office_admin', 'office_id': 'engineering-kano'})
    assert not dashboard.delete_history_btn.isVisible()
    assert not dashboard.delete_selected_btn.isVisible()

    # 4. Verify DeleteHistoryDialog criteria and params
    emps = {str(kano): {'_id': kano, 'employee_no': 'K1', 'full_name': 'Kano Staff', 'office_id': 'engineering-kano'}}
    dialog = DeleteHistoryDialog(employees=emps)
    qtbot.addWidget(dialog)

    assert dialog.radio_staff.isChecked()
    assert not dialog.delete_btn.isEnabled()
    dialog.confirm_input.setText("DELETE")
    assert dialog.delete_btn.isEnabled()

    params = dialog.get_delete_params()
    assert params.get('employee_id') == str(kano)
    assert 'work_date' in params


def test_dashboard_excel_and_csv_export(qtbot, tmp_path):
    from src.ui.dashboard import DashboardWidget
    dashboard = DashboardWidget()
    qtbot.addWidget(dashboard)
    dashboard.add_row(dashboard.attendance_table, [
        'Kano', 'Engineering Kano', 'Test User', 'T100', 'Engineer', '2026-10-21', '08:00', '17:00', 'Present'
    ])
    xlsx_path = str(tmp_path / 'report.xlsx')
    dashboard._export_to_excel(xlsx_path)
    assert os.path.exists(xlsx_path)

    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb['Attendance Sheet']
    assert ws['A1'].value == 'RANDAFRAMES ATTENDANCE REPORT'
    assert ws['C5'].value == 'Test User'

    csv_path = str(tmp_path / 'report.csv')
    dashboard._export_to_csv(csv_path)
    assert os.path.exists(csv_path)
    with open(csv_path, 'r', encoding='utf-8-sig') as f:
        content = f.read()
    assert 'Test User' in content


def test_office_admin_dashboard_date_grouping_and_gaps(qtbot, monkeypatch, tmp_path):
    from src.ui.dashboard import DashboardWidget
    kano = ObjectId()

    def call(path, body=None):
        if body['collection'] == 'employees':
            return [{'office_id': 'engineering-kano', '_id': kano, 'employee_no': 'K1', 'full_name': 'Kano Staff', 'status': 'active'}]
        # Multiple dates for Kano office
        return [
            {'office_id': 'engineering-kano', 'employee_id': kano, 'work_date': '2026-10-21', 'clock_in_at': datetime(2026, 10, 21, 8)},
            {'office_id': 'engineering-kano', 'employee_id': kano, 'work_date': '2026-10-20', 'clock_in_at': datetime(2026, 10, 20, 8, 15)}
        ]

    monkeypatch.setattr(api, 'call', call)
    monkeypatch.setattr(api, 'token', 'test-token')

    dashboard = DashboardWidget()
    qtbot.addWidget(dashboard)

    # Set as Office Admin for Kano
    dashboard.set_user({'role': 'office_admin', 'office_id': 'engineering-kano'})
    assert dashboard.title.text() == 'Engineering Kano'

    table = dashboard.attendance_table
    # Expected structure:
    # Row 0: Date Header (2026-10-21) - span 9 (no emoji)
    # Row 1: Data row for 2026-10-21 - span 1
    # Row 2: Spacer Gap row - span 9
    # Row 3: Date Header (2026-10-20) - span 9
    # Row 4: Data row for 2026-10-20 - span 1
    assert table.rowCount() == 5
    assert table.columnCount() == 9
    assert table.columnSpan(0, 0) == 9
    assert 'Date:' in table.item(0, 0).text()
    assert '📅' not in table.item(0, 0).text()
    assert '2026' in table.item(0, 0).text()

    assert table.columnSpan(1, 0) == 1
    assert table.item(1, 2).text() == 'Kano Staff'

    assert table.columnSpan(2, 0) == 9  # Gap row
    assert table.item(2, 0).text() == ''

    assert table.columnSpan(3, 0) == 9
    assert 'Date:' in table.item(3, 0).text()

    assert table.columnSpan(4, 0) == 1
    assert table.item(4, 2).text() == 'Kano Staff'

    # Verify export includes date section headers and data records
    csv_path = str(tmp_path / 'office_report.csv')
    dashboard._export_to_csv(csv_path)
    with open(csv_path, 'r', encoding='utf-8-sig') as f:
        lines = [line.strip() for line in f if line.strip()]
    # 1 column header + 2 date headers + 2 data rows = 5 lines
    assert len(lines) == 5
    assert 'Date:' in lines[1]
    assert 'Date:' in lines[3]

    # Verify Excel export includes date section headers
    xlsx_path = str(tmp_path / 'office_report.xlsx')
    dashboard._export_to_excel(xlsx_path)
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb['Attendance Sheet']
    assert 'Date:' in str(ws['A5'].value)


def test_password_line_edit_visibility_toggle(qtbot):
    from src.ui.password_input import PasswordLineEdit
    from PySide6.QtWidgets import QLineEdit

    pw_input = PasswordLineEdit("Test Password")
    qtbot.addWidget(pw_input)

    assert pw_input.echoMode() == QLineEdit.Password
    assert "Show password" in pw_input.eye_action.toolTip()

    # Trigger visibility toggle
    pw_input.toggle_visibility()
    assert pw_input.echoMode() == QLineEdit.Normal
    assert "Hide password" in pw_input.eye_action.toolTip()

    # Toggle back
    pw_input.toggle_visibility()
    assert pw_input.echoMode() == QLineEdit.Password
    assert "Show password" in pw_input.eye_action.toolTip()


def test_theme_manager_and_toggle_button(qtbot):
    from src.ui.theme import ThemeManager, ThemeToggleButton

    tm = ThemeManager.instance()
    initial_dark = tm.is_dark

    btn = ThemeToggleButton()
    qtbot.addWidget(btn)

    # Click toggle button
    btn.click()
    assert tm.is_dark != initial_dark

    # Click again to restore
    btn.click()
    assert tm.is_dark == initial_dark


def test_scanner_worker_hold_and_best_frame(monkeypatch):
    import src.scanner.worker as worker_mod
    from src.scanner.worker import ScannerWorker
    from unittest.mock import MagicMock

    # Patch hold duration to 0.05s and sleep to instant for fast test execution
    monkeypatch.setattr(worker_mod, 'HOLD_DURATION', 0.05)
    monkeypatch.setattr(worker_mod.time, 'sleep', lambda s: None)

    mock_scanner = MagicMock()
    mock_scanner.get_device_count.return_value = 1
    mock_scanner._handle = 1
    mock_scanner.begin_capture.return_value = True

    # High contrast frame (alternating 0 and 255) vs lower contrast (all 128)
    frame_low = b'\x80' * 100
    frame_high = (b'\x00\xff' * 50)

    frames = [
        (True, frame_low, 10, 10, False, 0),
        (True, frame_high, 10, 10, False, 0),
        (True, frame_low, 10, 10, False, 0),
    ]

    # Advance time slightly on each frame so elapsed >= HOLD_DURATION is reached
    fake_clock = [1000.0, 1000.02, 1000.07, 1000.10]
    monkeypatch.setattr(worker_mod.time, 'time', lambda: fake_clock.pop(0) if fake_clock else 1001.0)

    def mock_get_image(ignore_preview=False):
        if frames:
            return frames.pop(0)
        return (True, frame_low, 10, 10, False, 0)

    mock_scanner.block_get_image.side_effect = mock_get_image

    worker = ScannerWorker(mock_scanner)
    completed = []
    statuses = []

    worker.scan_complete.connect(lambda data: completed.append(data))
    worker.scan_status.connect(lambda msg: statuses.append(msg))

    worker.run()

    assert len(completed) == 1
    assert completed[0]['status'] == 'success'
    # High-contrast frame had the highest sharpness/std dev and was selected!
    assert completed[0]['image'] == frame_high
    assert any('Scanning: Hold finger steady' in s for s in statuses)


def test_offices_tab_visibility_and_table(qtbot, monkeypatch):
    from src.ui.dashboard import DashboardWidget
    from PySide6.QtWidgets import QApplication

    def dummy_call(path, body=None):
        if path == '/read':
            if body.get('collection') == 'offices':
                return [{'location': 'Kano', 'name': 'Engineering Kano', '_id': 'engineering-kano'}]
            if body.get('collection') == 'employees':
                return []
            if body.get('collection') == 'attendance_records':
                return []
        return []

    monkeypatch.setattr(api, 'call', dummy_call)
    monkeypatch.setattr(api, 'token', 'test-token')

    dashboard = DashboardWidget()
    qtbot.addWidget(dashboard)

    # Verify tab names and order: 'Add offices' must be immediately after 'Manage employees'
    assert dashboard.tabs.tabText(0) == 'Attendance sheet'
    assert dashboard.tabs.tabText(1) == 'Manage employees'
    assert dashboard.tabs.tabText(2) == 'Add offices'

    # Master Admin sees the Add offices tab
    dashboard.set_user({'role': 'main_admin'})
    assert dashboard.tabs.isTabVisible(2) is True
    assert dashboard.offices_table.rowCount() >= 1

    # Office Admin cannot see the Add offices tab
    dashboard.set_user({'role': 'office_admin', 'office_id': 'engineering-kano'})
    assert dashboard.tabs.isTabVisible(2) is False


def test_add_office_dialog_password_and_clipboard(qtbot, monkeypatch):
    from src.ui.add_office import AddOfficeDialog
    from PySide6.QtWidgets import QApplication, QMessageBox

    created_payloads = []

    def mock_call(path, body=None):
        if path == '/offices/create':
            created_payloads.append(body)
            return {
                'ok': True,
                'office': {'_id': body['office_id'], 'name': body['name'], 'location': body['location']},
                'username': body['username'],
                'password': body['password']
            }
        return {}

    monkeypatch.setattr(api, 'call', mock_call)
    monkeypatch.setattr(QMessageBox, 'information', lambda *args, **kwargs: None)

    dialog = AddOfficeDialog()
    qtbot.addWidget(dialog)

    # 1. Password must be automatically generated
    initial_pw = dialog.pw_input.text()
    assert len(initial_pw) >= 10

    # 2. Clicking Generate Password creates a new password and copies it to clipboard
    dialog._regenerate_and_copy_password()
    new_pw = dialog.pw_input.text()
    assert new_pw != initial_pw
    assert QApplication.clipboard().text() == new_pw

    # 3. Copy Password button copies the password to clipboard
    dialog.pw_input.setText("CustomP@ssw0rd!9")
    dialog._copy_password_only()
    assert QApplication.clipboard().text() == "CustomP@ssw0rd!9"

    # 4. Fill in office details and verify slug generation
    dialog.name_input.setText("Logistics Asaba")
    assert dialog.id_input.text() == "logistics-asaba"
    assert dialog.username_input.text() == "logistics-asaba"
    dialog.location_input.setText("Asaba")

    # 5. Copy all credentials button formats and copies to clipboard
    dialog._copy_credentials()
    clip_text = QApplication.clipboard().text()
    assert "Office Name: Logistics Asaba" in clip_text
    assert "Location: Asaba" in clip_text
    assert "Office ID: logistics-asaba" in clip_text
    assert "Username: logistics-asaba" in clip_text
    assert "Temporary Password: CustomP@ssw0rd!9" in clip_text

    # 6. Submit creation: calls API, registers office, copies to clipboard
    dialog._on_create()
    assert len(created_payloads) == 1
    assert created_payloads[0]['name'] == 'Logistics Asaba'
    assert created_payloads[0]['office_id'] == 'logistics-asaba'
    assert created_payloads[0]['username'] == 'logistics-asaba'
    assert created_payloads[0]['password'] == 'CustomP@ssw0rd!9'

    from src.offices import OFFICE_MAP
    assert 'logistics-asaba' in OFFICE_MAP


def test_manual_clock_and_digest_dialogs_instantiate(qtbot):
    from src.ui.manual_clock import ManualClockInDialog
    from src.ui.digest_settings import DigestSettingsDialog

    manual_dialog = ManualClockInDialog(employees={}, parent=None)
    qtbot.addWidget(manual_dialog)
    assert manual_dialog.windowTitle() == "Supervisor Manual Staff Clock-In / Out"

    digest_dialog = DigestSettingsDialog(user={"role": "main_admin"}, parent=None)
    qtbot.addWidget(digest_dialog)
    assert digest_dialog.windowTitle() == "Daily Morning Email Digest Settings"
