import csv
from datetime import datetime, timedelta, timezone
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                             QTabWidget, QComboBox, QDateEdit, QLineEdit, QTableWidget,
                             QTableWidgetItem, QHeaderView, QMessageBox, QFileDialog, QInputDialog,
                             QDialog, QRadioButton, QButtonGroup, QGroupBox)
from PySide6.QtCore import Qt, QDate, Signal, QThread, QTimer
from PySide6.QtGui import QGuiApplication, QColor, QFont
from bson import ObjectId
from src.client import api, ClientDatabase
from src.offices import OFFICES, OFFICE_MAP

LAGOS = timezone(timedelta(hours=1))


def local_time(value):
    if not value:
        return '--:--'
    return value.replace(tzinfo=timezone.utc).astimezone(LAGOS).strftime('%H:%M:%S')


def csv_safe(value):
    text = str(value or '')
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else text


class DashboardLoader(QThread):
    data_loaded = Signal(list, list, int)
    load_failed = Signal(str)

    def __init__(self, emp_query, att_query):
        super().__init__()
        self.emp_query = emp_query
        self.att_query = att_query

    def run(self):
        try:
            db = ClientDatabase.get_db()
            employees = list(db.employees.find(self.emp_query))
            records = list(db.attendance_records.find(self.att_query))
            today = datetime.now(LAGOS).strftime('%Y-%m-%d')
            present_count = len({r['employee_id'] for r in records if r.get('work_date') == today and r.get('clock_in_at')})
            self.data_loaded.emit(employees, records, present_count)
        except Exception as exc:
            self.load_failed.emit(str(exc))


class PasswordResetWorker(QThread):
    reset_success = Signal(dict, str)
    reset_failed = Signal(str)

    def __init__(self, office_id, office_name):
        super().__init__()
        self.office_id = office_id
        self.office_name = office_name

    def run(self):
        try:
            result = api.call(f'/offices/{self.office_id}/reset-password')
            self.reset_success.emit(result, self.office_name)
        except Exception as exc:
            self.reset_failed.emit(str(exc))


class PasswordChangeWorker(QThread):
    change_success = Signal()
    change_failed = Signal(str)

    def __init__(self, password):
        super().__init__()
        self.password = password

    def run(self):
        try:
            api.call('/password', {'password': self.password})
            self.change_success.emit()
        except Exception as exc:
            self.change_failed.emit(str(exc))


class DeactivateWorker(QThread):
    deactivate_success = Signal()
    deactivate_failed = Signal(str)

    def __init__(self, employee_id):
        super().__init__()
        self.employee_id = employee_id

    def run(self):
        try:
            ClientDatabase.get_db().employees.update_one(
                {'_id': self.employee_id},
                {'$set': {'status': 'inactive'}}
            )
            self.deactivate_success.emit()
        except Exception as exc:
            self.deactivate_failed.emit(str(exc))


class DeleteAttendanceWorker(QThread):
    finished_success = Signal(int)
    finished_error = Signal(str)

    def __init__(self, record_id=None, employee_id=None, work_date=None, before_date=None, from_date=None, to_date=None, office_id=None, delete_all=False):
        super().__init__()
        self.record_id = record_id
        self.employee_id = employee_id
        self.work_date = work_date
        self.before_date = before_date
        self.from_date = from_date
        self.to_date = to_date
        self.office_id = office_id
        self.delete_all = delete_all

    def run(self):
        try:
            res = api.delete_attendance(
                record_id=self.record_id,
                employee_id=self.employee_id,
                work_date=self.work_date,
                before_date=self.before_date,
                from_date=self.from_date,
                to_date=self.to_date,
                office_id=self.office_id,
                delete_all=self.delete_all
            )
            self.finished_success.emit(res.get("deleted_count", 0))
        except Exception as exc:
            self.finished_error.emit(str(exc))


class DeleteHistoryDialog(QDialog):
    def __init__(self, employees=None, parent=None):
        if isinstance(employees, QWidget) and parent is None:
            parent = employees
            employees = None
        super().__init__(parent)
        self.employees = employees or {}
        self.setWindowTitle("Delete Attendance History")
        self.setFixedWidth(520)
        layout = QVBoxLayout(self)
        layout.setSpacing(12)

        title = QLabel("Delete Attendance History")
        title.setStyleSheet("font-size: 16px; font-weight: bold; color: #ff6b6b;")
        layout.addWidget(title)

        desc = QLabel("Permanently remove attendance records from the central database. Restricted to Master Administrator.")
        desc.setWordWrap(True)
        desc.setStyleSheet("color: #aaa; font-size: 12px; margin-bottom: 4px;")
        layout.addWidget(desc)

        group = QGroupBox("Deletion Target")
        g_layout = QVBoxLayout(group)
        g_layout.setSpacing(10)

        self.btn_group = QButtonGroup(self)
        self.radio_staff = QRadioButton("Particular staff log on a specific day")
        self.radio_range = QRadioButton("Records within date range")
        self.radio_all = QRadioButton("ALL attendance records across all time")

        self.btn_group.addButton(self.radio_staff, 1)
        self.btn_group.addButton(self.radio_range, 2)
        self.btn_group.addButton(self.radio_all, 3)

        today = QDate.currentDate()

        g_layout.addWidget(self.radio_staff)
        h_staff = QHBoxLayout()
        h_staff.setContentsMargins(24, 0, 0, 4)
        self.staff_combo = QComboBox()
        self.staff_combo.setFixedWidth(250)
        sorted_emps = sorted(self.employees.values(), key=lambda e: (e.get('office_id', ''), e.get('full_name', '')))
        for emp in sorted_emps:
            name = emp.get('full_name') or f"{emp.get('first_name', '')} {emp.get('last_name', '')}".strip()
            office = emp.get('office_id', '')
            emp_no = emp.get('employee_no', '')
            label = f"{name} ({emp_no}) - {office}"
            self.staff_combo.addItem(label, str(emp.get('_id', '')))
        if not sorted_emps:
            self.staff_combo.addItem("No staff records loaded", "")

        self.staff_date_edit = QDateEdit(today)
        self.staff_date_edit.setCalendarPopup(True)
        self.staff_date_edit.setFixedWidth(120)
        h_staff.addWidget(self.staff_combo)
        h_staff.addWidget(QLabel("Date:"))
        h_staff.addWidget(self.staff_date_edit)
        h_staff.addStretch()
        g_layout.addLayout(h_staff)

        g_layout.addWidget(self.radio_range)
        self.from_date_edit = QDateEdit(today.addDays(-60))
        self.from_date_edit.setCalendarPopup(True)
        self.from_date_edit.setFixedWidth(120)
        self.to_date_edit = QDateEdit(today.addDays(-30))
        self.to_date_edit.setCalendarPopup(True)
        self.to_date_edit.setFixedWidth(120)
        h_range = QHBoxLayout()
        h_range.setContentsMargins(24, 0, 0, 4)
        h_range.addWidget(QLabel("From:"))
        h_range.addWidget(self.from_date_edit)
        h_range.addWidget(QLabel("To:"))
        h_range.addWidget(self.to_date_edit)
        h_range.addStretch()
        g_layout.addLayout(h_range)

        g_layout.addWidget(self.radio_all)
        layout.addWidget(group)

        self.radio_staff.setChecked(True)
        self.radio_staff.toggled.connect(self._update_radio_states)
        self.radio_range.toggled.connect(self._update_radio_states)
        self.radio_all.toggled.connect(self._update_radio_states)
        self._update_radio_states()

        h_office = QHBoxLayout()
        h_office.addWidget(QLabel("Target Office Filter:"))
        self.office_combo = QComboBox()
        self.office_combo.addItem("All Offices (Entire Company)", "")
        for o in OFFICES:
            self.office_combo.addItem(o["name"], o["_id"])
        h_office.addWidget(self.office_combo)
        layout.addLayout(h_office)

        confirm_box = QVBoxLayout()
        confirm_label = QLabel("Type DELETE to confirm permanent deletion:")
        confirm_label.setStyleSheet("font-weight: bold; color: #ff8888; font-size: 12px;")
        confirm_box.addWidget(confirm_label)

        self.confirm_input = QLineEdit()
        self.confirm_input.setPlaceholderText("DELETE")
        self.confirm_input.setStyleSheet("padding: 6px; font-weight: bold; font-size: 13px;")
        confirm_box.addWidget(self.confirm_input)
        layout.addLayout(confirm_box)

        btn_layout = QHBoxLayout()
        btn_layout.addStretch()

        self.cancel_btn = QPushButton("Cancel")
        self.cancel_btn.setFixedHeight(34)
        self.cancel_btn.clicked.connect(self.reject)
        btn_layout.addWidget(self.cancel_btn)

        self.delete_btn = QPushButton("Permanently Delete")
        self.delete_btn.setFixedHeight(34)
        self.delete_btn.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #8a1d29, stop:1 #ad2637);
                color: #ffffff;
                border: 1px solid #c93b50;
                border-radius: 5px;
                padding: 0 20px;
                font-weight: bold;
                font-size: 13px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #a82333, stop:1 #c72b41);
                border: 1px solid #ff576f;
            }
            QPushButton:pressed {
                background-color: #63121d;
                border: 1px solid #d9384e;
            }
            QPushButton:disabled {
                background-color: #2b1d20;
                color: #6b5256;
                border-color: #3d282c;
            }
        """)
        self.delete_btn.setEnabled(False)
        self.delete_btn.clicked.connect(self.accept)
        btn_layout.addWidget(self.delete_btn)
        layout.addLayout(btn_layout)

        self.confirm_input.textChanged.connect(self._check_confirm)
        self.staff_combo.currentIndexChanged.connect(lambda: self._check_confirm(self.confirm_input.text()))

    def _update_radio_states(self):
        staff_active = self.radio_staff.isChecked()
        self.staff_combo.setEnabled(staff_active)
        self.staff_date_edit.setEnabled(staff_active)

        range_active = self.radio_range.isChecked()
        self.from_date_edit.setEnabled(range_active)
        self.to_date_edit.setEnabled(range_active)
        if hasattr(self, 'confirm_input'):
            self._check_confirm(self.confirm_input.text())

    def _check_confirm(self, text):
        confirmed = text.strip() == "DELETE"
        if self.radio_staff.isChecked() and not self.staff_combo.currentData():
            confirmed = False
        self.delete_btn.setEnabled(confirmed)

    def get_delete_params(self):
        office_id = self.office_combo.currentData() or None
        if self.radio_staff.isChecked():
            emp_id = self.staff_combo.currentData() or None
            work_date = self.staff_date_edit.date().toString("yyyy-MM-dd")
            return {"employee_id": emp_id, "work_date": work_date, "office_id": office_id}
        elif self.radio_all.isChecked():
            return {"delete_all": True, "office_id": office_id}
        else:
            from_str = self.from_date_edit.date().toString("yyyy-MM-dd")
            to_str = self.to_date_edit.date().toString("yyyy-MM-dd")
            return {"from_date": from_str, "to_date": to_str, "office_id": office_id}


class DashboardWidget(QWidget):
    logout_requested = Signal()
    kiosk_requested = Signal()

    def __init__(self, scanner_provider=None, parent=None):
        super().__init__(parent)
        self.setObjectName("dashboardWidget")
        self.setAttribute(Qt.WA_StyledBackground, True)
        self.scanner_provider = scanner_provider
        self.user = None
        self.employees = {}
        self.loader = None
        self.inactivity_timer = QTimer(self)
        self.inactivity_timer.setInterval(600_000)
        self.inactivity_timer.timeout.connect(self.on_inactivity_timeout)
        layout = QVBoxLayout(self)

        header = QHBoxLayout()
        self.title = QLabel('Randaframes Attendance')
        self.title.setStyleSheet('font-size: 22px; font-weight: bold;')
        header.addWidget(self.title)
        header.addStretch()

        from src.ui.theme import ThemeToggleButton, ThemeManager
        self.theme_btn = ThemeToggleButton()
        header.addWidget(self.theme_btn)
        ThemeManager.instance().theme_changed.connect(self.on_theme_changed)

        self.kiosk_button = QPushButton('Open attendance terminal')
        self.kiosk_button.clicked.connect(self.kiosk_requested.emit)
        header.addWidget(self.kiosk_button)

        password_button = QPushButton('Change password')
        password_button.clicked.connect(self.change_password)
        header.addWidget(password_button)

        logout = QPushButton('Log out')
        logout.clicked.connect(self.logout_requested.emit)
        header.addWidget(logout)
        layout.addLayout(header)

        self.summary = QLabel('')
        self.summary.setStyleSheet('font-size: 14px; font-weight: bold; color: #4CAF50; padding: 6px 0;')
        layout.addWidget(self.summary)

        filters = QHBoxLayout()
        self.location_filter = QComboBox()
        self.location_filter.addItem('All locations', '')
        for loc in sorted({o['location'] for o in OFFICES}):
            self.location_filter.addItem(loc, loc)
        filters.addWidget(QLabel('Location:'))
        filters.addWidget(self.location_filter)

        self.office_filter = QComboBox()
        self.office_filter.addItem('All offices', '')
        for o in OFFICES:
            self.office_filter.addItem(o['name'], o['_id'])
        filters.addWidget(QLabel('Office:'))
        filters.addWidget(self.office_filter)

        today = QDate.currentDate()
        self.from_date = QDateEdit(today.addDays(-7))
        self.from_date.setCalendarPopup(True)
        self.to_date = QDateEdit(today)
        self.to_date.setCalendarPopup(True)
        filters.addWidget(QLabel('From:'))
        filters.addWidget(self.from_date)
        filters.addWidget(QLabel('To:'))
        filters.addWidget(self.to_date)

        self.search = QLineEdit()
        self.search.setPlaceholderText('Search staff name or ID')
        self.search.textChanged.connect(self.apply_search_filter)
        filters.addWidget(self.search)

        refresh = QPushButton('Refresh')
        refresh.clicked.connect(self.refresh)
        filters.addWidget(refresh)

        export = QPushButton('Export Sheet')
        export.setToolTip('Export attendance records to Excel (.xlsx) or CSV')
        export.clicked.connect(self.export_reports)
        filters.addWidget(export)

        backups_btn = QPushButton('Daily Backups')
        backups_btn.setToolTip('Open the local SQLite daily attendance backups folder')
        backups_btn.clicked.connect(self.open_backups_folder)
        filters.addWidget(backups_btn)

        self.manual_clock_btn = QPushButton('Manual Staff Clock-in')
        self.manual_clock_btn.setToolTip('Supervisor Emergency Check-in: manually record attendance for cut/unreadable finger')
        self.manual_clock_btn.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #1e40af, stop:1 #2563eb);
                color: #ffffff;
                font-weight: 600;
                font-size: 12px;
                border: 1px solid #3b82f6;
                border-radius: 5px;
                padding: 5px 12px;
            }
            QPushButton:hover {
                background: #1d4ed8;
            }
        """)
        self.manual_clock_btn.clicked.connect(self.open_manual_clock_in)
        filters.addWidget(self.manual_clock_btn)

        self.digest_btn = QPushButton('Email Digest')
        self.digest_btn.setToolTip('Configure recipient emails and preview the daily morning attendance digest')
        self.digest_btn.clicked.connect(self.open_digest_settings)
        filters.addWidget(self.digest_btn)

        self.delete_selected_btn = QPushButton('Delete Selected Log')
        self.delete_selected_btn.setToolTip('Permanently delete the selected attendance entry (Master Admin only)')
        self.delete_selected_btn.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #3d141b, stop:0.5 #541b25, stop:1 #3d141b);
                color: #ff99a8;
                border: 1px solid #7c2431;
                border-radius: 5px;
                font-size: 12px;
                font-weight: 600;
                padding: 5px 14px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #5e1c28, stop:0.5 #7a2233, stop:1 #5e1c28);
                color: #ffffff;
                border: 1px solid #ff4d6d;
            }
            QPushButton:pressed {
                background-color: #240c11;
                border: 1px solid #c92a43;
            }
            QPushButton:disabled {
                background: #1e191b;
                color: #554448;
                border: 1px solid #2d2225;
            }
        """)
        self.delete_selected_btn.setEnabled(False)
        self.delete_selected_btn.setVisible(False)
        self.delete_selected_btn.clicked.connect(self.delete_selected_log)
        filters.addWidget(self.delete_selected_btn)
        layout.addLayout(filters)

        self.tabs = QTabWidget()
        layout.addWidget(self.tabs)

        self.attendance_table = self.table(['Location', 'Office', 'Name', 'Employee ID',
                                           'Job Title', 'Date', 'Time In', 'Time Out', 'Status'])
        self.attendance_table.itemSelectionChanged.connect(self.on_attendance_selection_changed)
        self.tabs.addTab(self.attendance_table, 'Attendance sheet')

        employees = QWidget()
        employee_layout = QVBoxLayout(employees)
        emp_buttons = QHBoxLayout()

        add_emp = QPushButton('Register staff')
        add_emp.clicked.connect(self.open_add_employee)
        emp_buttons.addWidget(add_emp)

        edit_emp = QPushButton('Edit staff')
        edit_emp.clicked.connect(self.edit_employee)
        emp_buttons.addWidget(edit_emp)

        enroll_btn = QPushButton('Enroll fingerprint')
        enroll_btn.clicked.connect(self.enroll)
        emp_buttons.addWidget(enroll_btn)

        deactivate_btn = QPushButton('Deactivate staff')
        deactivate_btn.clicked.connect(self.deactivate)
        emp_buttons.addWidget(deactivate_btn)

        emp_buttons.addStretch()
        employee_layout.addLayout(emp_buttons)

        self.emp_table = self.table(['Office', 'Employee ID', 'Name', 'Job Title', 'Status', 'Fingerprint'])
        employee_layout.addWidget(self.emp_table)
        self.tabs.addTab(employees, 'Manage employees')

        # Offices Tab (Master Admin only) - Named 'Add offices' next to 'Manage employees'
        self.offices_widget = QWidget()
        offices_layout = QVBoxLayout(self.offices_widget)
        offices_layout.setContentsMargins(0, 10, 0, 0)
        offices_layout.setSpacing(10)

        office_actions = QHBoxLayout()
        self.add_office_btn = QPushButton('+ Add New Office')
        self.add_office_btn.setFixedHeight(34)
        self.add_office_btn.setStyleSheet("""
            QPushButton {
                background-color: #2563eb;
                color: #ffffff;
                border-radius: 6px;
                font-size: 13px;
                font-weight: 600;
                padding: 0 16px;
            }
            QPushButton:hover {
                background-color: #1d4ed8;
            }
            QPushButton:pressed {
                background-color: #1e40af;
            }
        """)
        self.add_office_btn.clicked.connect(self.open_add_office)
        office_actions.addWidget(self.add_office_btn)

        self.reset_office_pw_btn = QPushButton('Reset Office Password')
        self.reset_office_pw_btn.setFixedHeight(34)
        self.reset_office_pw_btn.setToolTip('Reset temporary password for the selected office and copy to clipboard')
        self.reset_office_pw_btn.clicked.connect(self.reset_selected_office_password)
        office_actions.addWidget(self.reset_office_pw_btn)

        self.delete_office_btn = QPushButton('Delete Office')
        self.delete_office_btn.setFixedHeight(34)
        self.delete_office_btn.setToolTip('Delete the selected branch office and revoke its admin credentials (core offices are protected)')
        self.delete_office_btn.setStyleSheet("""
            QPushButton {
                background-color: transparent;
                color: #ef4444;
                border: 1px solid #ef4444;
                border-radius: 6px;
                font-size: 13px;
                font-weight: 500;
                padding: 0 14px;
            }
            QPushButton:hover {
                background-color: #ef4444;
                color: #ffffff;
            }
        """)
        self.delete_office_btn.clicked.connect(self.delete_selected_office)
        office_actions.addWidget(self.delete_office_btn)

        self.refresh_offices_btn = QPushButton('Refresh')
        self.refresh_offices_btn.setFixedHeight(34)
        self.refresh_offices_btn.clicked.connect(self.populate_offices_table)
        office_actions.addWidget(self.refresh_offices_btn)

        office_actions.addStretch()
        offices_layout.addLayout(office_actions)

        self.offices_table = self.table(['Location', 'Office Name', 'Office ID', 'Admin Username', 'Status', 'Staff Count'])
        offices_layout.addWidget(self.offices_table)
        self.tabs.addTab(self.offices_widget, 'Add offices')

        master_actions = QHBoxLayout()
        master_actions.setContentsMargins(0, 10, 0, 4)
        master_actions.setSpacing(12)
        master_actions.addStretch()

        self.reset_button = QPushButton('Reset Office Password')
        self.reset_button.setFixedHeight(34)
        self.reset_button.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #1e293b, stop:1 #0f172a);
                color: #94a3b8;
                border: 1px solid #334155;
                border-radius: 6px;
                font-size: 12px;
                font-weight: 600;
                padding: 0 16px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #334155, stop:1 #1e293b);
                color: #e2e8f0;
                border: 1px solid #64748b;
            }
            QPushButton:pressed {
                background-color: #0b1120;
                border-color: #475569;
            }
        """)
        self.reset_button.clicked.connect(self.reset_office_password)
        master_actions.addWidget(self.reset_button)

        self.delete_history_btn = QPushButton('Delete History...')
        self.delete_history_btn.setFixedHeight(34)
        self.delete_history_btn.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #2b1116, stop:0.5 #3d141e, stop:1 #2b1116);
                color: #ff99a8;
                border: 1px solid #6b1f2b;
                border-radius: 6px;
                font-size: 12px;
                font-weight: 600;
                padding: 0 18px;
                letter-spacing: 0.3px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #421620, stop:0.5 #5e1c2b, stop:1 #421620);
                color: #ffffff;
                border: 1px solid #ff4d6d;
            }
            QPushButton:pressed {
                background-color: #1a080c;
                border: 1px solid #a82338;
            }
            QPushButton:disabled {
                background: #181214;
                color: #554448;
                border-color: #2b1c20;
            }
        """)
        self.delete_history_btn.clicked.connect(self.open_delete_history_dialog)
        master_actions.addWidget(self.delete_history_btn)
        layout.addLayout(master_actions)

        self.location_filter.currentIndexChanged.connect(self.refresh)
        self.office_filter.currentIndexChanged.connect(self.refresh)
        self.cached_records = []

    @staticmethod
    def table(headers):
        table = QTableWidget(0, len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QTableWidget.NoEditTriggers)
        table.setSelectionBehavior(QTableWidget.SelectRows)
        table.setSelectionMode(QTableWidget.SingleSelection)
        table.setAlternatingRowColors(True)
        return table

    def set_user(self, user):
        self.user = user
        self.inactivity_timer.start(600_000)
        master = user['role'] == 'main_admin'
        self.title.setText('Master Admin (All Offices)' if master else OFFICE_MAP[user['office_id']]['name'])
        self.kiosk_button.setVisible(not master)
        self.reset_button.setVisible(master)
        self.delete_history_btn.setVisible(master)
        self.delete_selected_btn.setVisible(master)
        self.delete_selected_btn.setEnabled(False)
        self.location_filter.setEnabled(master)
        self.office_filter.setEnabled(master)
        self.location_filter.blockSignals(True)
        self.office_filter.blockSignals(True)
        self.location_filter.setCurrentIndex(0)
        self.office_filter.setCurrentIndex(0 if master else self.office_filter.findData(user['office_id']))
        self.location_filter.blockSignals(False)
        self.office_filter.blockSignals(False)
        if hasattr(self, 'offices_widget'):
            offices_idx = self.tabs.indexOf(self.offices_widget)
            if offices_idx != -1:
                self.tabs.setTabVisible(offices_idx, master)
        if master:
            self.populate_offices_table()
        self.refresh()

    def clear_data(self):
        self.inactivity_timer.stop()
        self.user = None
        self.employees = {}
        self.cached_records = []
        self.delete_selected_btn.setEnabled(False)
        self.emp_table.setRowCount(0)
        self.attendance_table.setRowCount(0)
        if hasattr(self, 'offices_table'):
            self.offices_table.setRowCount(0)

    def reset_inactivity_timer(self):
        if self.user:
            self.inactivity_timer.start(600_000)

    def on_inactivity_timeout(self):
        self.inactivity_timer.stop()
        if self.user:
            QMessageBox.warning(self, 'Session Expired', 'Your admin session has expired after 10 minutes of inactivity for security.')
            self.logout_requested.emit()

    def mousePressEvent(self, event):
        self.reset_inactivity_timer()
        super().mousePressEvent(event)

    def keyPressEvent(self, event):
        self.reset_inactivity_timer()
        super().keyPressEvent(event)

    def query(self):
        if self.user and self.user['role'] == 'office_admin':
            return {'office_id': self.user['office_id']}
        selected = self.office_filter.currentData()
        location = self.location_filter.currentData()
        ids = [o['_id'] for o in OFFICES if (not selected or o['_id'] == selected)
               and (not location or o['location'] == location)]
        return {'office_id': {'$in': ids}}

    def refresh(self, sync=None):
        if not self.user or not api.token:
            return
        if self.from_date.date() > self.to_date.date():
            QMessageBox.warning(self, 'Invalid Date', 'The start date must be before the end date.')
            return

        import os
        is_sync = sync if sync is not None else (os.environ.get('QT_QPA_PLATFORM') == 'offscreen')
        emp_query = self.query()
        att_query = dict(self.query(), work_date={'$gte': self.from_date.date().toString('yyyy-MM-dd'),
                                                  '$lte': self.to_date.date().toString('yyyy-MM-dd')})

        if is_sync:
            try:
                db = ClientDatabase.get_db()
                employees = list(db.employees.find(emp_query))
                records = list(db.attendance_records.find(att_query))
                today = datetime.now(LAGOS).strftime('%Y-%m-%d')
                present_count = len({r['employee_id'] for r in records if r.get('work_date') == today and r.get('clock_in_at')})
                self.on_data_loaded(employees, records, present_count)
            except Exception as exc:
                self.on_load_failed(str(exc))
            return

        self.summary.setText('Loading records from server...')

        if self.loader and self.loader.isRunning():
            self.loader.terminate()
            self.loader.wait()

        self.loader = DashboardLoader(emp_query, att_query)
        self.loader.data_loaded.connect(self.on_data_loaded)
        self.loader.load_failed.connect(self.on_load_failed)
        self.loader.start()

    def on_data_loaded(self, employees, records, present_count):
        self.emp_table.setRowCount(0)
        self.employees = {str(e['_id']): e for e in employees}
        self.cached_records = records

        for employee in sorted(employees, key=lambda e: (e.get('office_id', ''), e.get('full_name', ''))):
            values = [OFFICE_MAP.get(employee.get('office_id'), {}).get('name', 'Unassigned'),
                      employee.get('employee_no'), employee.get('full_name'), employee.get('job_title'),
                      employee.get('status'), 'Enrolled' if employee.get('has_fingerprint') else 'Needs enrollment']
            self.add_row(self.emp_table, values, str(employee['_id']))

        self.populate_attendance_table(records)

        if self.user and self.user.get('role') == 'main_admin' and hasattr(self, 'offices_table'):
            self.populate_offices_table()

        active_ids = {e['_id'] for e in employees if e.get('status') == 'active'}
        self.summary.setText(f'Active staff: {len(active_ids)}   |   Checked in today: {present_count}'
                             f'   |   Not checked in: {max(0, len(active_ids) - present_count)}'
                             f'   |   Displayed records: {len(records)}')

        try:
            from src.services.backup import BackupService
            today_str = datetime.now(LAGOS).strftime('%Y-%m-%d')
            today_records = [r for r in records if r.get('work_date') == today_str]
            if today_records:
                BackupService().backup_daily_attendance(records=today_records, employees=self.employees)
        except Exception:
            pass

    def populate_attendance_table(self, records):
        self.attendance_table.setRowCount(0)
        self.attendance_table.clearSpans()
        search = self.search.text().strip().lower()

        filtered_records = []
        for record in records:
            employee = self.employees.get(str(record.get('employee_id')), {})
            name = employee.get('full_name', 'Archived employee')
            number = employee.get('employee_no', '')
            if search and search not in (name + ' ' + number).lower():
                continue
            filtered_records.append(record)

        if not filtered_records:
            return

        from collections import defaultdict
        grouped = defaultdict(list)
        for record in filtered_records:
            d = record.get('work_date') or 'Unscheduled'
            grouped[d].append(record)

        sorted_dates = sorted(grouped.keys(), reverse=True)

        for date_idx, work_date in enumerate(sorted_dates):
            date_records = grouped[work_date]
            date_records.sort(key=lambda r: (
                OFFICE_MAP.get(r.get('office_id'), {}).get('location', ''),
                OFFICE_MAP.get(r.get('office_id'), {}).get('name', ''),
                self.employees.get(str(r.get('employee_id')), {}).get('full_name', '')
            ))

            try:
                dt = datetime.strptime(work_date, "%Y-%m-%d")
                date_display = dt.strftime("%A, %d %B %Y  (%d/%m/%Y)")
            except Exception:
                date_display = work_date

            present_count = sum(1 for r in date_records if r.get('clock_in_at'))
            header_text = f"  Date: {date_display}    •    {len(date_records)} record(s)    •    {present_count} checked in"

            # 1. Section Header Row for this Date
            h_row = self.attendance_table.rowCount()
            self.attendance_table.insertRow(h_row)
            self.attendance_table.setRowHeight(h_row, 36)
            self.attendance_table.setSpan(h_row, 0, 1, self.attendance_table.columnCount())

            from src.ui.theme import ThemeManager
            is_dark = ThemeManager.instance().is_dark
            h_bg = QColor("#1e293b") if is_dark else QColor("#e0e7ff")
            h_fg = QColor("#38bdf8") if is_dark else QColor("#1e3a8a")
            gap_bg = QColor("#000000") if is_dark else QColor("#f3f4f6")

            h_item = QTableWidgetItem(header_text)
            h_item.setFlags(Qt.ItemIsEnabled)
            h_item.setBackground(h_bg)
            h_item.setForeground(h_fg)
            h_font = QFont("Segoe UI", 11, QFont.Bold)
            h_item.setFont(h_font)
            self.attendance_table.setItem(h_row, 0, h_item)

            for col in range(1, self.attendance_table.columnCount()):
                dummy = QTableWidgetItem("")
                dummy.setFlags(Qt.NoItemFlags)
                dummy.setBackground(h_bg)
                self.attendance_table.setItem(h_row, col, dummy)

            # 2. Staff Attendance Rows for this Date
            for record in date_records:
                employee = self.employees.get(str(record.get('employee_id')), {})
                name = employee.get('full_name', 'Archived employee')
                number = employee.get('employee_no', '')
                office = OFFICE_MAP.get(record.get('office_id'), {})
                values = [
                    office.get('location', 'Unassigned'),
                    office.get('name', 'Unassigned'),
                    name,
                    number,
                    employee.get('job_title', ''),
                    record.get('work_date', ''),
                    local_time(record.get('clock_in_at')),
                    local_time(record.get('clock_out_at')),
                    record.get('status', '')
                ]
                rec_id = str(record.get('_id', ''))
                self.add_row(self.attendance_table, values, record_id=rec_id)
                curr_row = self.attendance_table.rowCount() - 1
                self.attendance_table.setRowHeight(curr_row, 30)

            # 3. Gap / Spacer Row after each Date group (except the last date)
            if date_idx < len(sorted_dates) - 1:
                gap_row = self.attendance_table.rowCount()
                self.attendance_table.insertRow(gap_row)
                self.attendance_table.setRowHeight(gap_row, 16)
                self.attendance_table.setSpan(gap_row, 0, 1, self.attendance_table.columnCount())

                gap_item = QTableWidgetItem("")
                gap_item.setFlags(Qt.NoItemFlags)
                gap_item.setBackground(gap_bg)
                self.attendance_table.setItem(gap_row, 0, gap_item)
                for col in range(1, self.attendance_table.columnCount()):
                    d_gap = QTableWidgetItem("")
                    d_gap.setFlags(Qt.NoItemFlags)
                    d_gap.setBackground(gap_bg)
                    self.attendance_table.setItem(gap_row, col, d_gap)

    def apply_search_filter(self):
        if hasattr(self, 'cached_records'):
            self.populate_attendance_table(self.cached_records)

    def on_theme_changed(self, is_dark):
        try:
            import shiboken6
            if not shiboken6.isValid(self) or not shiboken6.isValid(self.attendance_table):
                return
            self.apply_search_filter()
        except Exception:
            pass

    def on_load_failed(self, error_msg):
        self.employees = {}
        self.emp_table.setRowCount(0)
        self.attendance_table.setRowCount(0)
        self.summary.setText('Could not load records. Check connection and refresh.')
        QMessageBox.warning(self, 'Refresh Failed', error_msg)

    @staticmethod
    def add_row(table, values, record_id=None):
        row = table.rowCount()
        table.insertRow(row)
        for column, value in enumerate(values):
            item = QTableWidgetItem('' if value is None else str(value))
            if column == 0:
                item.setData(Qt.UserRole, record_id)
            table.setItem(row, column, item)

    def selected_employee(self):
        row = self.emp_table.currentRow()
        if row < 0:
            QMessageBox.information(self, 'Select employee', 'Select an employee first.')
            return None
        return self.employees.get(self.emp_table.item(row, 0).data(Qt.UserRole))

    def open_add_employee(self):
        from src.ui.add_employee import AddEmployeeDialog
        dialog = AddEmployeeDialog(self)
        if dialog.exec():
            self.refresh()

    def edit_employee(self):
        from src.ui.add_employee import AddEmployeeDialog
        employee = self.selected_employee()
        if employee and AddEmployeeDialog(self, employee).exec():
            self.refresh()

    def enroll(self):
        employee = self.selected_employee()
        if employee:
            from src.ui.enrollment import EnrollFingerprintDialog
            EnrollFingerprintDialog(str(employee['_id']), self.scanner_provider, self).exec()
            self.refresh()

    def deactivate(self):
        employee = self.selected_employee()
        if not employee:
            return
        if QMessageBox.question(self, 'Deactivate employee', 'Prevent future check-ins? Attendance history will be retained.') != QMessageBox.Yes:
            return

        self.summary.setText('Deactivating employee...')
        self.deact_worker = DeactivateWorker(employee['_id'])
        self.deact_worker.deactivate_success.connect(self.on_deactivate_success)
        self.deact_worker.deactivate_failed.connect(self.on_deactivate_failed)
        self.deact_worker.start()

    def on_deactivate_success(self):
        self.summary.setText('Employee deactivated successfully.')
        self.refresh()

    def on_deactivate_failed(self, error_msg):
        self.summary.setText('Could not deactivate employee.')
        QMessageBox.warning(self, 'Could not deactivate', error_msg)

    def export_reports(self):
        default_name = f"Attendance_Report_{self.from_date.date().toString('yyyyMMdd')}_{self.to_date.date().toString('yyyyMMdd')}.xlsx"
        path, selected_filter = QFileDialog.getSaveFileName(
            self, 'Export sheet', default_name, 'Excel Workbook (*.xlsx);;CSV File (*.csv)'
        )
        if not path:
            return
        try:
            if path.lower().endswith('.xlsx') or ('xlsx' in selected_filter.lower() and not path.lower().endswith('.csv')):
                if not path.lower().endswith('.xlsx'):
                    path += '.xlsx'
                self._export_to_excel(path)
            else:
                self._export_to_csv(path)
            QMessageBox.information(self, 'Export Complete', f'Attendance records successfully exported to:\n\n{path}')
        except Exception as exc:
            QMessageBox.warning(self, 'Export Failed', str(exc))

    def _export_to_excel(self, path):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Sheet"
        ws.views.sheetView[0].showGridLines = True

        table = self.attendance_table
        col_count = table.columnCount()
        headers = [table.horizontalHeaderItem(c).text() for c in range(col_count)]

        end_col = get_column_letter(max(1, col_count))
        ws.merge_cells(f"A1:{end_col}1")
        title_cell = ws["A1"]
        title_cell.value = "RANDAFRAMES ATTENDANCE REPORT"
        title_cell.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 34

        office_name = self.office_filter.currentText() if hasattr(self, 'office_filter') else "All Offices"
        loc_name = self.location_filter.currentText() if hasattr(self, 'location_filter') else "All Locations"
        from_str = self.from_date.date().toString("yyyy-MM-dd")
        to_str = self.to_date.date().toString("yyyy-MM-dd")
        exported_at = datetime.now(LAGOS).strftime("%Y-%m-%d %H:%M:%S")

        meta_info = f"Office: {office_name}   |   Location: {loc_name}   |   Period: {from_str} to {to_str}   |   Generated: {exported_at}"
        ws.merge_cells(f"A2:{end_col}2")
        meta_cell = ws["A2"]
        meta_cell.value = meta_info
        meta_cell.font = Font(name="Arial", size=9, italic=True, color="475569")
        meta_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 8

        header_row = 4
        ws.row_dimensions[header_row].height = 26
        header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for col_idx, header_text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        data_font = Font(name="Arial", size=10)

        status_col_idx = None
        for i, h in enumerate(headers, 1):
            if "status" in h.lower():
                status_col_idx = i

        present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        present_font = Font(name="Arial", size=10, bold=True, color="166534")

        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        absent_font = Font(name="Arial", size=10, bold=True, color="991B1B")

        date_banner_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        date_banner_font = Font(name="Arial", size=11, bold=True, color="38BDF8")
        date_banner_border = Border(
            left=Side(style="thin", color="334155"),
            right=Side(style="thin", color="334155"),
            top=Side(style="medium", color="0284C7"),
            bottom=Side(style="medium", color="0284C7")
        )

        excel_row = header_row
        for row_idx in range(table.rowCount()):
            if table.columnSpan(row_idx, 0) > 1:
                item0 = table.item(row_idx, 0)
                text0 = item0.text().strip() if item0 else ""
                excel_row += 1
                if text0:
                    ws.merge_cells(f"A{excel_row}:{end_col}{excel_row}")
                    cell = ws.cell(row=excel_row, column=1)
                    cell.value = f"   {text0}"
                    cell.font = date_banner_font
                    cell.fill = date_banner_fill
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    ws.row_dimensions[excel_row].height = 24
                    for c_idx in range(1, col_count + 1):
                        ws.cell(row=excel_row, column=c_idx).border = date_banner_border
                else:
                    ws.merge_cells(f"A{excel_row}:{end_col}{excel_row}")
                    ws.row_dimensions[excel_row].height = 10
                continue

            excel_row += 1
            ws.row_dimensions[excel_row].height = 20
            row_fill = zebra_fill if (excel_row % 2 == 1) else white_fill

            for col_idx in range(1, col_count + 1):
                raw_item = table.item(row_idx, col_idx - 1)
                val = raw_item.text().strip() if raw_item else ""
                cell = ws.cell(row=excel_row, column=col_idx)

                cell.value = val
                if col_idx in (1, 2, 3, 5):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                cell.font = data_font
                cell.fill = row_fill
                cell.border = thin_border

                if col_idx == status_col_idx:
                    status_lower = val.lower()
                    if "present" in status_lower:
                        cell.fill = present_fill
                        cell.font = present_font
                    elif "absent" in status_lower or "not checked" in status_lower:
                        cell.fill = absent_fill
                        cell.font = absent_font

        for col_idx in range(1, col_count + 1):
            max_len = len(headers[col_idx - 1])
            for r in range(header_row + 1, excel_row + 1):
                val_str = str(ws.cell(row=r, column=col_idx).value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(13, min(max_len + 4, 38))

        ws.freeze_panes = f"A{header_row + 1}"
        wb.save(path)

    def _export_to_csv(self, path):
        with open(path, 'w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.writer(handle)
            table = self.attendance_table
            writer.writerow([table.horizontalHeaderItem(c).text() for c in range(table.columnCount())])
            for row in range(table.rowCount()):
                if table.columnSpan(row, 0) > 1:
                    item0 = table.item(row, 0)
                    text0 = item0.text().strip() if item0 else ""
                    if text0:
                        writer.writerow([text0] + [''] * (table.columnCount() - 1))
                    else:
                        writer.writerow([])
                    continue
                writer.writerow([csv_safe(table.item(row, col).text() if table.item(row, col) else '') for col in range(table.columnCount())])

    def open_backups_folder(self):
        from src.services.backup import BackupService
        service = BackupService()
        if not service.open_backup_folder():
            QMessageBox.information(self, 'Backups Location', f'Your daily SQLite backups are saved in:\n\n{service.backup_dir}')

    def populate_offices_table(self):
        if not hasattr(self, 'offices_table'):
            return
        from src.offices import OFFICES, sync_offices
        try:
            db_offices = list(ClientDatabase.get_db().offices.find({}))
            sync_offices(db_offices)
        except Exception:
            pass

        self.offices_table.setRowCount(0)
        for office in sorted(OFFICES, key=lambda o: (o.get('location', ''), o.get('name', ''))):
            oid = office.get('_id', '')
            name = office.get('name', '')
            loc = office.get('location', '')
            admin_user = office.get('username') or oid
            if self.employees:
                staff_count = sum(1 for e in self.employees.values() if e.get('office_id') == oid)
            else:
                try:
                    staff_count = ClientDatabase.get_db().employees.count_documents({'office_id': oid})
                except Exception:
                    staff_count = 0
            values = [loc, name, oid, admin_user, "Active", str(staff_count)]
            self.add_row(self.offices_table, values, record_id=oid)

    def open_add_office(self):
        from src.ui.add_office import AddOfficeDialog
        dialog = AddOfficeDialog(self)
        if dialog.exec():
            self._update_office_combos()
            self.populate_offices_table()
            self.refresh()

    def _update_office_combos(self):
        from src.offices import OFFICES
        curr_loc = self.location_filter.currentData()
        curr_off = self.office_filter.currentData()
        self.location_filter.blockSignals(True)
        self.office_filter.blockSignals(True)

        self.location_filter.clear()
        self.location_filter.addItem('All locations', '')
        for loc in sorted({o['location'] for o in OFFICES if o.get('location')}):
            self.location_filter.addItem(loc, loc)

        self.office_filter.clear()
        self.office_filter.addItem('All offices', '')
        for o in OFFICES:
            self.office_filter.addItem(o['name'], o['_id'])

        idx_l = self.location_filter.findData(curr_loc)
        if idx_l >= 0:
            self.location_filter.setCurrentIndex(idx_l)
        idx_o = self.office_filter.findData(curr_off)
        if idx_o >= 0:
            self.office_filter.setCurrentIndex(idx_o)

        self.location_filter.blockSignals(False)
        self.office_filter.blockSignals(False)

    def reset_selected_office_password(self):
        row = self.offices_table.currentRow()
        if row < 0:
            return self.reset_office_password()

        item0 = self.offices_table.item(row, 0)
        office_id = item0.data(Qt.UserRole) if item0 else None
        if not office_id and self.offices_table.item(row, 2):
            office_id = self.offices_table.item(row, 2).text().strip()

        office_name = self.offices_table.item(row, 1).text() if self.offices_table.item(row, 1) else (office_id or "Selected Office")
        if not office_id:
            return self.reset_office_password()

        if QMessageBox.question(
            self,
            'Reset Office Password',
            f'This will generate a new temporary password for "{office_name}" and sign out all of its active terminals.\n\nContinue?'
        ) != QMessageBox.Yes:
            return

        self.summary.setText(f'Resetting password for {office_name}...')
        self.reset_worker = PasswordResetWorker(office_id, office_name)
        self.reset_worker.reset_success.connect(self.on_reset_success)
        self.reset_worker.reset_failed.connect(self.on_reset_failed)
        self.reset_worker.start()

    def delete_selected_office(self):
        row = self.offices_table.currentRow()
        if row < 0:
            QMessageBox.warning(self, "No Office Selected", "Please select an office row from the table to delete.")
            return

        item0 = self.offices_table.item(row, 0)
        office_id = item0.data(Qt.UserRole) if item0 else None
        if not office_id and self.offices_table.item(row, 2):
            office_id = self.offices_table.item(row, 2).text().strip()

        office_name = self.offices_table.item(row, 1).text() if self.offices_table.item(row, 1) else (office_id or "Selected Office")
        if not office_id:
            QMessageBox.warning(self, "Error", "Could not determine the selected office ID.")
            return

        from src.offices import CORE_OFFICE_IDS
        if office_id in CORE_OFFICE_IDS:
            QMessageBox.warning(
                self,
                "Protected Office",
                f"The core default office '{office_name}' is protected and cannot be deleted."
            )
            return

        staff_item = self.offices_table.item(row, 5)
        staff_count = 0
        if staff_item:
            try:
                staff_count = int(staff_item.text().strip())
            except ValueError:
                staff_count = 0

        if staff_count > 0:
            QMessageBox.warning(
                self,
                "Cannot Delete Office",
                f"Office '{office_name}' currently has {staff_count} active employee(s) assigned.\n\n"
                f"Please reassign or deactivate the staff members before deleting this office."
            )
            return

        reply = QMessageBox.question(
            self,
            "Confirm Delete Office",
            f"Are you sure you want to delete office '{office_name}' ({office_id})?\n\n"
            f"• The office will be removed from all active dropdowns.\n"
            f"• Its administrator credentials will be revoked immediately.\n"
            f"• All past attendance records remain preserved for historical reports.\n\n"
            f"Proceed with deletion?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )
        if reply != QMessageBox.Yes:
            return

        try:
            api.call(f"/offices/{office_id}/delete")
            from src.offices import unregister_office
            unregister_office(office_id)
            if hasattr(self, '_update_office_combos'):
                self._update_office_combos()
            self.populate_offices_table()
            self.refresh()
            QMessageBox.information(self, "Office Deleted", f"Office '{office_name}' has been successfully deleted.")
        except Exception as exc:
            QMessageBox.critical(self, "Delete Failed", str(exc))

    def open_manual_clock_in(self):
        from src.ui.manual_clock import ManualClockInDialog
        dialog = ManualClockInDialog(employees=self.employees, parent=self)
        if dialog.exec():
            self.refresh()

    def open_digest_settings(self):
        from src.ui.digest_settings import DigestSettingsDialog
        dialog = DigestSettingsDialog(user=self.user, parent=self)
        dialog.exec()

    def reset_office_password(self):
        office = self.office_filter.currentData()
        if not office:
            items = [OFFICE_MAP[oid]['name'] for oid in OFFICE_MAP]
            item, ok = QInputDialog.getItem(self, 'Reset Office Password', 'Select Office Account:', items, 0, False)
            if not ok or not item:
                return
            office = next(oid for oid in OFFICE_MAP if OFFICE_MAP[oid]['name'] == item)

        office_name = OFFICE_MAP.get(office, {}).get('name', office)
        if QMessageBox.question(self, 'Reset Office Password',
                                f'This will generate a new temporary password for {office_name} and sign out all of its active terminals. Continue?') != QMessageBox.Yes:
            return

        self.summary.setText(f'Resetting password for {office_name}...')
        self.reset_worker = PasswordResetWorker(office, office_name)
        self.reset_worker.reset_success.connect(self.on_reset_success)
        self.reset_worker.reset_failed.connect(self.on_reset_failed)
        self.reset_worker.start()

    def on_reset_success(self, result, office_name):
        self.summary.setText('Password reset successfully.')
        QGuiApplication.clipboard().setText(result['password'])
        QMessageBox.information(
            self,
            'Password Reset Complete',
            f"Office: {office_name}\n"
            f"Username: {result['username']}\n"
            f"New Temporary Password: {result['password']}\n\n"
            f"(Copied to clipboard. Deliver to office admin securely. They will be prompted to change it at next login.)"
        )

    def on_reset_failed(self, error_msg):
        self.summary.setText('Password reset failed.')
        QMessageBox.warning(self, 'Reset Failed', error_msg)

    def change_password(self):
        from src.ui.login import ChangePasswordDialog
        dialog = ChangePasswordDialog(
            title="Change Password",
            message="Set new password (at least 10 characters with letters, numbers, and symbols):",
            parent=self
        )
        if dialog.exec() == QDialog.Accepted and dialog.get_password():
            password = dialog.get_password()
            self.summary.setText('Updating your password...')
            self.pw_worker = PasswordChangeWorker(password)
            self.pw_worker.change_success.connect(self.on_change_password_success)
            self.pw_worker.change_failed.connect(self.on_change_password_failed)
            self.pw_worker.start()

    def on_change_password_success(self):
        self.summary.setText('Password updated successfully.')
        QMessageBox.information(self, 'Password updated', 'Your password was changed.')

    def on_change_password_failed(self, error_msg):
        self.summary.setText('Password change failed.')
        QMessageBox.warning(self, 'Password not changed', error_msg)

    def on_attendance_selection_changed(self):
        is_master = bool(self.user and self.user.get('role') == 'main_admin')
        has_sel = bool(self.attendance_table.selectionModel() and self.attendance_table.selectionModel().hasSelection())
        if has_sel and is_master:
            row = self.attendance_table.currentRow()
            if row >= 0 and self.attendance_table.columnSpan(row, 0) == 1:
                self.delete_selected_btn.setEnabled(True)
                return
        self.delete_selected_btn.setEnabled(False)

    def delete_selected_log(self):
        if not self.user or self.user.get('role') != 'main_admin':
            return
        row = self.attendance_table.currentRow()
        if row < 0 or self.attendance_table.columnSpan(row, 0) > 1:
            QMessageBox.information(self, 'Select Log', 'Select an attendance record from the table first.')
            return

        item0 = self.attendance_table.item(row, 0)
        record_id = item0.data(Qt.UserRole) if item0 else None
        location = self.attendance_table.item(row, 0).text() if self.attendance_table.item(row, 0) else ''
        office = self.attendance_table.item(row, 1).text() if self.attendance_table.item(row, 1) else ''
        staff_name = self.attendance_table.item(row, 2).text() if self.attendance_table.item(row, 2) else ''
        emp_no = self.attendance_table.item(row, 3).text() if self.attendance_table.item(row, 3) else ''
        work_date = self.attendance_table.item(row, 5).text() if self.attendance_table.item(row, 5) else ''
        time_in = self.attendance_table.item(row, 6).text() if self.attendance_table.item(row, 6) else ''

        emp_id = None
        for eid, edata in self.employees.items():
            if edata.get('employee_no') == emp_no or edata.get('full_name') == staff_name:
                emp_id = eid
                break

        if not record_id and not (emp_id and work_date):
            QMessageBox.information(self, 'Select Log', 'Select an attendance record from the table first.')
            return

        msg = (f"Are you sure you want to permanently delete this attendance record?\n\n"
               f"• Staff: {staff_name} ({emp_no})\n"
               f"• Office: {office} ({location})\n"
               f"• Date: {work_date}\n"
               f"• Time In: {time_in or 'N/A'}\n\n"
               f"This action is immediate and cannot be undone.")

        if QMessageBox.question(self, 'Confirm Deletion', msg, QMessageBox.Yes | QMessageBox.No, QMessageBox.No) != QMessageBox.Yes:
            return

        self.reset_inactivity_timer()
        self.delete_selected_btn.setEnabled(False)
        self.summary.setText(f"Deleting record for {staff_name} on {work_date}...")

        self.del_worker = DeleteAttendanceWorker(
            record_id=record_id or None,
            employee_id=emp_id if not record_id else None,
            work_date=work_date if not record_id else None
        )
        self.del_worker.finished_success.connect(self.on_delete_history_success)
        self.del_worker.finished_error.connect(self.on_delete_history_error)
        self.del_worker.start()

    def open_delete_history_dialog(self):
        if not self.user or self.user.get('role') != 'main_admin':
            return
        self.reset_inactivity_timer()
        dialog = DeleteHistoryDialog(employees=self.employees, parent=self)
        if dialog.exec():
            params = dialog.get_delete_params()
            self.delete_history_btn.setEnabled(False)
            self.delete_history_btn.setText("Deleting...")
            self.summary.setText("Deleting attendance history...")
            self.del_worker = DeleteAttendanceWorker(**params)
            self.del_worker.finished_success.connect(self.on_delete_history_success)
            self.del_worker.finished_error.connect(self.on_delete_history_error)
            self.del_worker.start()

    def on_delete_history_success(self, count):
        self.delete_history_btn.setEnabled(True)
        self.delete_history_btn.setText("Delete History...")
        self.on_attendance_selection_changed()
        self.summary.setText(f"Deleted {count} record(s).")
        QMessageBox.information(self, "Success", f"Successfully deleted {count} attendance record(s).")
        self.refresh()

    def on_delete_history_error(self, err_msg):
        self.delete_history_btn.setEnabled(True)
        self.delete_history_btn.setText("Delete History...")
        self.on_attendance_selection_changed()
        self.summary.setText("Attendance deletion failed.")
        QMessageBox.critical(self, "Deletion Failed", err_msg)

